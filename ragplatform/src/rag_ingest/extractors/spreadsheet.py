"""
XLSX and CSV/TSV (spec stage 3): one `sheet_rows` element per block of rows.

Each element is a Markdown table that repeats the sheet's header row, and
records the header and the exact spreadsheet row numbers in its attributes, so
the sheet chunker can regroup rows without losing either. Empty sheets are
skipped; very large sheets are capped with a warning rather than silently cut.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
from pathlib import PurePosixPath
from typing import Any, Iterable

from rag_ingest.config import Settings
from rag_ingest.extractors.base import decode_text
from rag_ingest.ids import element_id
from rag_ingest.models import DocumentModel, Element, Provenance
from rag_ingest.pipeline.errors import PermanentError

NAME, VERSION = "spreadsheet", "1"

LEGACY = {".xls", ".xlt", ".ods", ".ots"}


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value).replace("\r\n", " ").replace("\n", " ").strip()


def markdown_row(cells: list[str]) -> str:
    return "| " + " | ".join(c.replace("|", "\\|") for c in cells) + " |"


def markdown_table(header: list[str], rows: list[list[str]]) -> str:
    lines = [markdown_row(header), "| " + " | ".join("---" for _ in header) + " |"]
    lines.extend(markdown_row(row) for row in rows)
    return "\n".join(lines)


def _sheets_xlsx(raw: bytes) -> Iterable[tuple[str, Iterable[tuple[int, list[str]]]]]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001 — corrupt or encrypted workbook
        raise PermanentError("extraction_failed", f"openpyxl could not read the workbook: {error}") from error

    def rows(sheet):
        for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            yield number, [cell_text(v) for v in row]

    try:
        for sheet in book.worksheets:
            yield sheet.title, rows(sheet)
    finally:
        book.close()


def _sheets_csv(raw: bytes, suffix: str) -> Iterable[tuple[str, Iterable[tuple[int, list[str]]]]]:
    text = decode_text(raw)
    if suffix == ".tsv":
        dialect: Any = csv.excel_tab
    else:
        try:
            dialect = csv.Sniffer().sniff(text[:20_000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    yield "", ((number, [c.strip() for c in row]) for number, row in enumerate(reader, start=1))


def extract(document_id: str, content_hash: str, content_type: str, raw: bytes, filename: str,
            settings: Settings) -> DocumentModel:
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix in LEGACY:
        raise PermanentError("unsupported_format",
                             f"{suffix} workbooks are not supported; save the file as .xlsx")
    sheets = _sheets_csv(raw, suffix) if suffix in {".csv", ".tsv"} or not raw.startswith(b"PK") \
        else _sheets_xlsx(raw)

    elements: list[Element] = []
    warnings: list[str] = []
    summaries: list[dict[str, Any]] = []
    per_element = max(1, settings.sheet_rows_per_element)

    for name, rows in sheets:
        header: list[str] | None = None
        block: list[tuple[int, list[str]]] = []
        count = 0
        truncated = False

        def flush() -> None:
            if not block or header is None:
                return
            first, last = block[0][0], block[-1][0]
            elements.append(Element(
                id=element_id(document_id, len(elements)),
                type="sheet_rows",
                # A wide row may have widened the header after earlier rows were read.
                text=markdown_table(header, [cells + [""] * (len(header) - len(cells)) for _, cells in block]),
                provenance=Provenance(sheet=name or None, row_range=(first, last)),
                attributes={"header": header, "row_numbers": [n for n, _ in block]},
            ))
            block.clear()

        for number, cells in rows:
            while cells and not cells[-1]:
                cells.pop()
            if not any(cells):
                continue
            if header is None:
                header = cells
                continue
            if count >= settings.max_sheet_rows:
                truncated = True
                break
            width = len(header)
            if len(cells) > width:
                header = header + [f"column {i + 1}" for i in range(width, len(cells))]
                width = len(header)
            block.append((number, cells + [""] * (width - len(cells))))
            count += 1
            if len(block) >= per_element:
                flush()
        flush()

        if header is None:
            continue                                  # an empty sheet
        if truncated:
            warnings.append(f"sheet {name or 'csv'}: only the first {settings.max_sheet_rows} rows were indexed")
        summaries.append({"sheet": name or None, "columns": header, "rows": count})

    if not elements and not summaries:
        raise PermanentError("low_quality_extraction", "the spreadsheet has no non-empty rows")
    return DocumentModel(
        document_id=document_id, content_hash=content_hash, content_type=content_type,
        content_family="spreadsheet",
        metadata={"sheets": summaries},
        elements=elements,
        extraction={"extractor": NAME, "version": VERSION, "warnings": warnings},
    )
