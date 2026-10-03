"""
Regenerates the binary fixtures for the Office / PDF / spreadsheet tests.

    python tests/fixtures/make_fixtures.py

The outputs are committed, so tests never depend on this script — it exists so
the fixtures can be read, reviewed and rebuilt rather than being opaque blobs.
Needs python-docx, python-pptx, openpyxl and Pillow (all pulled in by the
`docs` extra). The audio and video fixtures also need macOS `say` (for the
voices) and `ffmpeg`; on other systems that step is skipped.
"""

from __future__ import annotations

import csv
import shutil
import subprocess
import tempfile
import zlib
from pathlib import Path

HERE = Path(__file__).parent

CLAIM_ROWS = [
    ("CLM-1001", "Motor", "Open", "1200.00"),
    ("CLM-1002", "Home", "Closed", "5400.50"),
    ("CLM-1003", "Travel", "Open", "310.00"),
    ("CLM-1004", "Motor", "Referred", "18250.00"),
]


def docx() -> None:
    from docx import Document

    doc = Document()
    doc.core_properties.author = "Claims Operations"
    doc.add_heading("Claims Handling Guide", level=0)
    doc.add_heading("1 Intake", level=1)
    doc.add_paragraph(
        "Every first notice of loss is acknowledged within one business day. The handler "
        "records the policy number, the date of loss and a short description of what happened."
    )
    doc.add_heading("1.1 Required documents", level=2)
    for item in ("Completed claim form", "Photographs of the damage", "Police report for theft"):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("2 Settlement authority", level=1)
    doc.add_paragraph("Handlers may settle claims up to the limits below without referral.")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Grade", "Limit (GBP)", "Countersign")):
        cell.text = text
    # Enough rows that a small chunk budget must split the table.
    for grade in range(1, 41):
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = f"G{grade}", f"{grade * 2500:,}", "No" if grade < 30 else "Yes"
    doc.add_heading("3 Escalation", level=1)
    doc.add_paragraph("Claims above authority are referred to the duty manager with a recommendation.")
    doc.save(HERE / "claims-guide.docx")


def pptx() -> None:
    from pptx import Presentation

    deck = Presentation()
    slides = [
        ("Fraud indicators", "Late notification\nInconsistent accounts\nRecent policy changes",
         "Stress that no single indicator proves fraud."),
        ("Referral process", "Refer to the special investigations unit within 48 hours.",
         "SIU contact is on the intranet."),
        ("Questions", "", ""),
    ]
    for title, body, notes in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    deck.save(HERE / "fraud-training.pptx")


def xlsx_and_csv() -> None:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Claims"
    sheet.append(["Claim", "Line", "Status", "Reserve"])
    for row in CLAIM_ROWS * 15:          # 60 rows: several row groups
        sheet.append([row[0], row[1], row[2], float(row[3])])
    book.create_sheet("Empty")
    rates = book.create_sheet("Rates")
    rates.append(["Line", "Base rate"])
    rates.append(["Motor", 0.042])
    rates.append(["Home", 0.018])
    book.save(HERE / "claims-register.xlsx")

    with open(HERE / "reserves.csv", "w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Claim", "Line", "Status", "Reserve"])
        writer.writerows(CLAIM_ROWS)


def html() -> None:
    (HERE / "privacy-notice.html").write_text("""<!doctype html>
<html><head><title>Privacy notice</title></head>
<body>
<nav>Home | Products | Contact</nav>
<h1>Privacy notice</h1>
<p>We use your personal data to administer your policy and to handle claims.</p>
<h2>Retention</h2>
<table>
  <tr><th>Record</th><th>Kept for</th></tr>
  <tr><td>Policy documents</td><td>7 years</td></tr>
  <tr><td>Claim files</td><td>10 years</td></tr>
</table>
<h2>Your rights</h2>
<ul><li>Access your data</li><li>Correct your data</li><li>Ask us to delete it</li></ul>
</body></html>
""")


# ── PDFs, written by hand so no PDF library is needed ─────────────────────────

def _pdf(pages: list[bytes], images: dict[str, tuple[int, int, bytes]] | None = None) -> bytes:
    """Minimal PDF 1.4: Helvetica, one content stream per page, optional images."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    bold = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>")
    xobjects = {}
    for name, (width, height, rgb) in (images or {}).items():
        data = zlib.compress(rgb)
        xobjects[name] = add(
            b"<< /Type /XObject /Subtype /Image /Width %d /Height %d /ColorSpace /DeviceRGB "
            b"/BitsPerComponent 8 /Filter /FlateDecode /Length %d >>\nstream\n" % (width, height, len(data))
            + data + b"\nendstream")
    pages_id = len(objects) + 1 + 2 * len(pages)
    page_ids = []
    for content in pages:
        stream = add(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
        xo = b" ".join(b"/%s %d 0 R" % (n.encode(), i) for n, i in xobjects.items())
        page_ids.append(add(
            b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] /Contents %d 0 R "
            b"/Resources << /Font << /F1 %d 0 R /F2 %d 0 R >> /XObject << %s >> >> >>"
            % (pages_id, stream, font, bold, xo)))
    kids = b" ".join(b"%d 0 R" % i for i in page_ids)
    assert add(b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(page_ids))) == pages_id
    catalog = add(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id)

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, catalog, xref)
    return bytes(out)


def _text(x: int, y: int, size: int, text: str, bold: bool = False) -> bytes:
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return b"BT /%s %d Tf %d %d Td (%s) Tj ET\n" % (b"F2" if bold else b"F1", size, x, y, safe.encode("latin-1"))


def text_pdf() -> None:
    sections = [
        ("Section 1 Cover", [
            "We will pay for loss of or damage to your car caused by accident, fire or theft.",
            "The most we will pay is the market value of the car at the time of the loss."]),
        ("Section 2 Exclusions", [
            "We will not pay for wear and tear, mechanical failure or loss of use.",
            "We will not pay if the car was driven by someone not named on the certificate."]),
        ("Section 3 Claims", [
            "Tell us about any accident as soon as possible, and within seven days at the latest.",
            "Do not admit liability or offer payment without our written agreement."]),
    ]
    pages = []
    for number, (heading, paragraphs) in enumerate(sections, start=1):
        content = bytearray()
        content += _text(72, 760, 9, "Northwind Insurance - Motor policy wording")   # running header
        y = 700
        if number == 1:
            content += _text(72, y, 20, "Motor policy wording", bold=True)
            y -= 50
        content += _text(72, y, 14, heading, bold=True)
        y -= 30
        for paragraph in paragraphs:
            content += _text(72, y, 11, paragraph)
            y -= 24
        content += _text(280, 40, 9, f"Page {number} of {len(sections)}")             # running footer
        pages.append(bytes(content))
    (HERE / "motor-policy.pdf").write_bytes(_pdf(pages))


def _rendered_lines(lines: list[str], width: int = 1275, height: int = 1650):
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=40)
    y = 160
    for line in lines:
        draw.text((120, y), line, fill="black", font=font)
        y += 80
    return image


SCANNED_LINES = [
    "Witness statement",
    "I saw the blue van reverse into the parked car",
    "outside the bakery at about nine in the morning.",
    "The driver stopped and exchanged details.",
]


def scanned_pdf_and_png() -> None:
    image = _rendered_lines(SCANNED_LINES)
    image.save(HERE / "witness-statement.png")
    # A scan: the page is one image and has no text layer at all.
    width, height = image.size
    content = b"q 612 0 0 792 0 0 cm /Im1 Do Q\n"
    (HERE / "witness-statement-scan.pdf").write_bytes(
        _pdf([content], images={"Im1": (width, height, image.tobytes())}))


# ── Audio and video ───────────────────────────────────────────────────────────

CALL = [
    ("Samantha", "Good morning, you're through to the claims team. Can I take your policy number please?"),
    ("Daniel", "Yes, it's motor policy four four seven two. A van reversed into my parked car this morning."),
    ("Samantha", "I'm sorry to hear that. Was anyone hurt, and did you get the other driver's details?"),
    ("Daniel", "Nobody was hurt. The driver stopped and we exchanged details outside the bakery."),
    ("Samantha", "Thank you. I've opened claim C L M one zero zero five. "
                 "You'll receive an acknowledgement by email within one business day."),
]


def media() -> None:
    if not (shutil.which("say") and shutil.which("ffmpeg")):
        print("skipping audio/video fixtures: needs macOS `say` and ffmpeg")
        return
    with tempfile.TemporaryDirectory() as tmp:
        parts = []
        for i, (voice, line) in enumerate(CALL):
            part = Path(tmp) / f"{i}.aiff"
            subprocess.run(["say", "-v", voice, "-o", str(part), line], check=True)
            parts.append(part)
        # A second of silence between turns, as on a real call.
        silence = Path(tmp) / "silence.aiff"
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i",
                        "anullsrc=r=22050:cl=mono", "-t", "1", str(silence)], check=True)
        listing = Path(tmp) / "list.txt"
        listing.write_text("".join(f"file '{p}'\nfile '{silence}'\n" for p in parts))
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "concat", "-safe", "0", "-i", str(listing),
                        "-ac", "1", "-ar", "16000", "-b:a", "32k", str(HERE / "fnol-call.mp3")], check=True)
        # The same call as a video: a still frame over the audio. An explicit
        # duration, because -shortest overshoots with a 1 fps source.
        seconds = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0",
                                  str(HERE / "fnol-call.mp3")], check=True, capture_output=True, text=True).stdout.strip()
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "color=c=navy:s=320x240:r=1",
                        "-i", str(HERE / "fnol-call.mp3"), "-t", seconds, "-c:v", "libx264", "-tune", "stillimage",
                        "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "32k", str(HERE / "fnol-call.mp4")], check=True)


if __name__ == "__main__":
    docx()
    pptx()
    xlsx_and_csv()
    html()
    text_pdf()
    scanned_pdf_and_png()
    media()
    print("fixtures written to", HERE)
